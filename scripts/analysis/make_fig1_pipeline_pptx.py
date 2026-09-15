r"""
make_fig1_pipeline_pptx.py - Fig.1 v4 를 PowerPoint 네이티브 도형으로 재현.

모든 요소(장면·카드·캡슐·화살표·브래킷·텍스트)가 개별 편집 가능한 도형이다.
좌표는 make_fig1_pipeline.py v4 의 axes 좌표를 슬라이드 인치로 1:1 사상.
수학 기호는 유니코드(φ, μ, Σ, α)로 넣었다 — 아래첨자 등 세부는 PPT에서 직접 수정.

    python scripts/data/make_fig1_pipeline_pptx.py
"""
from __future__ import annotations
import os

import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(ROOT, "이미지", "fig1_pipeline.pptx")

SW, SH = 12.6, 3.43                       # 슬라이드 크기(in) = 그림 종횡비
E_GEN, E_FIX, E_LRN, E_CP = "1F5FA8", "BF7420", "2E7D4F", "A03A6B"
RED, INK, TAN = "C0392B", "141414", "8A6A25"

def C(h): return RGBColor.from_string(h)
def X(x): return Inches(x * SW)
def Y(y): return Inches((1.0 - y) * SH)   # axes(y-up) -> slide(y-down)
def W(w): return Inches(w * SW)
def H(h): return Inches(h * SH)
def FS(mpl): return Pt(round(mpl * 1.76)) # matplotlib pt -> 슬라이드 pt


def grad(shape, top, bot, edge, lw=1.6):
    shape.fill.gradient()
    st = shape.fill.gradient_stops
    st[0].color.rgb = C(top); st[0].position = 0.0
    st[1].color.rgb = C(bot); st[1].position = 1.0
    try:
        shape.fill.gradient_angle = 90.0
    except Exception:
        pass
    shape.line.color.rgb = C(edge)
    shape.line.width = Pt(lw)
    shape.shadow.inherit = False


def dash(line_obj):
    ln = line_obj._get_or_add_ln()
    d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
    ln.append(d)


def arrowhead(conn):
    ln = conn.line._get_or_add_ln()
    t = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(t)


def text(shapes, cx, cy, s, size, color=INK, bold=False, w=0.25, h=0.05,
         align=PP_ALIGN.CENTER, font="Arial"):
    tb = shapes.add_textbox(X(cx - w / 2), Y(cy + h / 2), W(w), H(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s
    r.font.size = size; r.font.bold = bold
    r.font.color.rgb = C(color); r.font.name = font
    return tb


def line(shapes, x1, y1, x2, y2, color, lw=1.6, dashed=False):
    cn = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(x1), Y(y1), X(x2), Y(y2))
    cn.line.color.rgb = C(color); cn.line.width = Pt(lw)
    cn.shadow.inherit = False
    if dashed:
        dash(cn.line)
    return cn


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sp = sl.shapes
    yc = 0.660

    # ---------- (1) 장면 ----------
    sx0, sx1, sy0, sy1 = 0.006, 0.192, 0.330, 0.985
    sc = sp.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(sx0), Y(sy1),
                      W(sx1 - sx0), H(sy1 - sy0))
    sc.adjustments[0] = 0.08
    grad(sc, "FDF6EA", "F6E3BD", "CAA25A", 1.4)
    scm = (sx0 + sx1) / 2
    text(sp, scm, 0.940, "synthetic scene  P0 / P1", FS(6.0), TAN, w=0.17)
    # ROI 점선 타원
    scx, scy, rx_r, ry_r = scm, 0.745, 0.064, 0.150
    ov = sp.add_shape(MSO_SHAPE.OVAL, X(scx - rx_r), Y(scy + ry_r),
                      W(2 * rx_r), H(2 * ry_r))
    ov.fill.background(); ov.line.color.rgb = C("CAA25A")
    ov.line.width = Pt(1.0); dash(ov.line); ov.shadow.inherit = False
    # 수신기 4대 + 방사원 + 레이
    rx_ang = [40, 140, 220, 320]
    rxs = [(scx + rx_r * np.cos(np.deg2rad(a)), scy + ry_r * np.sin(np.deg2rad(a)))
           for a in rx_ang]
    ex, ey = scx + 0.019, scy - 0.038
    for (px, py), cc in zip(rxs, [RED, RED, "3B6FB5", "3B6FB5"]):
        line(sp, px, py, ex, ey, cc, 0.9, dashed=True)
    for (px, py) in rxs:
        tr = sp.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                          X(px - 0.006), Y(py + 0.024), W(0.012), H(0.048))
        tr.fill.solid(); tr.fill.fore_color.rgb = C(E_GEN)
        tr.line.fill.background(); tr.shadow.inherit = False
        line(sp, px, py - 0.042, px, py - 0.012, E_GEN, 1.0)
    stx = sp.add_shape(MSO_SHAPE.STAR_5_POINT,
                       X(ex - 0.010), Y(ey + 0.036), W(0.020), H(0.072))
    stx.fill.solid(); stx.fill.fore_color.rgb = C("1A3A6B")
    stx.line.color.rgb = C("FFFFFF"); stx.line.width = Pt(0.5)
    stx.shadow.inherit = False
    text(sp, ex + 0.022, ey, "p", FS(7.0), "1A3A6B", w=0.03, font="Cambria Math")
    text(sp, scm, 0.520, "4 LPI/LPD waveforms", FS(5.3), TAN, w=0.17)
    text(sp, scm, 0.462, "(random - Gold-code DSSS - FH - burst)", FS(4.9), TAN, w=0.18)
    text(sp, scm, 0.392, "IQ:  TDOA imprint → TDL-D → AWGN", FS(5.1), TAN, w=0.18)

    # 장면 브래킷 (파랑) — 화살표 자리 갭
    ybl = 0.296
    for a, b in ((0.012, 0.036), (0.058, 0.186)):
        line(sp, a, ybl, b, ybl, E_GEN, 1.6)
    for xx in (0.012, 0.186):
        line(sp, xx, ybl, xx, ybl + 0.028, E_GEN, 1.6)
    text(sp, 0.130, 0.248, "labels p only under P0", FS(5.6), E_GEN, bold=True, w=0.16)

    # ---------- (2) GCC 카드 스택 ----------
    cx0, cy0, cw, chh = 0.262, 0.470, 0.104, 0.330
    for k in range(5, 0, -1):
        cd = sp.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                          X(cx0 - 0.0048 * k), Y(cy0 + chh + 0.021 * k),
                          W(cw), H(chh))
        cd.adjustments[0] = 0.10
        grad(cd, "F9DDB0", "E8A94F", E_FIX, 0.9)
    fc = sp.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(cx0), Y(cy0 + chh),
                      W(cw), H(chh))
    fc.adjustments[0] = 0.10
    grad(fc, "FFFDF8", "F6D59C", E_FIX, 1.4)
    # 앞 카드의 GCC 곡선 (freeform)
    xxw = np.linspace(0.04, 0.96, 60)
    ywv = (np.exp(-((xxw - 0.58) / 0.05) ** 2)
           + 0.25 * np.exp(-((xxw - 0.28) / 0.05) ** 2)
           + 0.10 * np.sin(29 * xxw) * np.exp(-((xxw - 0.5) / 0.33) ** 2))
    ywv = np.clip(ywv / ywv.max(), 0, None)
    pts = [(X(cx0 + cw * u), Y(cy0 + 0.052 + (chh - 0.125) * v))
           for u, v in zip(xxw, ywv)]
    fb = sp.build_freeform(pts[0][0], pts[0][1], scale=1)
    fb.add_line_segments(pts[1:], close=False)
    curve = fb.convert_to_shape()
    curve.fill.background(); curve.line.color.rgb = C(E_FIX)
    curve.line.width = Pt(1.4); curve.shadow.inherit = False
    text(sp, cx0 + cw / 2, cy0 + chh - 0.052, "GCC-PHAT", FS(6.6), INK,
         bold=True, w=0.10)
    text(sp, cx0 + cw / 2, cy0 - 0.055, "X : 6 × 257", FS(6.6), INK,
         bold=True, w=0.10, font="Cambria Math")
    text(sp, cx0 + cw / 2, cy0 - 0.115, "(6 pairs, ±128 lags)", FS(5.1),
         "6B6B6B", w=0.12)
    # fixed 브래킷
    fx0, fx1, ybf = 0.244, 0.378, 0.300
    line(sp, fx0, ybf, fx1, ybf, E_FIX, 1.6)
    for xx in (fx0, fx1):
        line(sp, xx, ybf, xx, ybf + 0.028, E_FIX, 1.6)
    text(sp, (fx0 + fx1) / 2, 0.252, "fixed - CRLB anchor", FS(5.6), E_FIX, w=0.14)

    # ---------- (3) 캡슐 ----------
    hh = 0.200

    def caps(x0, x1, label, edge, top, bot, fs=7.2):
        s = sp.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x0), Y(yc + hh / 2),
                         W(x1 - x0), H(hh))
        s.adjustments[0] = 0.42
        grad(s, top, bot, edge, 1.7)
        tf = s.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.size = FS(fs); r.font.bold = True
        r.font.color.rgb = C("1F1F1F"); r.font.name = "Arial"

    caps(0.408, 0.498, "1-D CNN", E_LRN, "EEF7F1", "B9DCC6")
    caps(0.552, 0.652, "NIW head", E_LRN, "EEF7F1", "B9DCC6")
    caps(0.706, 0.836, "split conformal", E_CP, "F8EBF2", "E3BCD2", 6.7)

    # ---------- (4) 출력 타원 ----------
    ox = 0.920
    ow, oh = 0.052, 0.26
    el = sp.add_shape(MSO_SHAPE.OVAL, X(ox - ow / 2), Y(yc + oh / 2),
                      W(ow), H(oh))
    el.rotation = -28
    el.fill.solid(); el.fill.fore_color.rgb = C("F6E6EE")
    el.line.color.rgb = C(E_CP); el.line.width = Pt(1.7)
    el.shadow.inherit = False
    dot = sp.add_shape(MSO_SHAPE.OVAL, X(ox - 0.004), Y(yc + 0.014),
                       W(0.008), H(0.028))
    dot.fill.solid(); dot.fill.fore_color.rgb = C("1A3A6B")
    dot.line.fill.background(); dot.shadow.inherit = False
    text(sp, ox, yc - 0.235, "Cα(X)", FS(6.6), INK, bold=True, w=0.08,
         font="Cambria Math")

    # ---------- (5) 체인 화살표 ----------
    def chain(xa, xb, lab=None):
        cn = line(sp, xa, yc, xb, yc, "595959", 1.7)
        arrowhead(cn)
        if lab:
            text(sp, (xa + xb) / 2, yc + 0.095, lab, FS(6.4), INK, bold=True,
                 w=0.10, font="Cambria Math")

    sx1_, cx0_ = 0.196, 0.232
    chain(sx1_, cx0_)
    chain(cx0 + cw + 0.005, 0.404)
    chain(0.502, 0.548, "φ (128)")
    chain(0.656, 0.702, "μ, Σa, Σe")
    chain(0.840, ox - 0.040)

    # ---------- (6) learned / frozen 브래킷 ----------
    def bracket(x0, x1, col, lab, ybr=0.470):
        line(sp, x0, ybr, x1, ybr, col, 1.6)
        for xx in (x0, x1):
            line(sp, xx, ybr, xx, ybr + 0.028, col, 1.6)
        text(sp, (x0 + x1) / 2, ybr - 0.070, lab, FS(5.6), col, w=0.22)

    bracket(0.408, 0.652, E_LRN, "learned")
    bracket(0.706, 0.836, E_CP, "calibrated on P0, then frozen")

    # ---------- (7) 하단 3패널 ----------
    RY0, RY1 = 0.028, 0.192

    def panel(x0, x1, edge, title, sub, fs_sub=5.6):
        pn = sp.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x0), Y(RY1),
                          W(x1 - x0), H(RY1 - RY0))
        pn.adjustments[0] = 0.10
        pn.fill.solid(); pn.fill.fore_color.rgb = C("FFFFFF")
        pn.line.color.rgb = C(edge); pn.line.width = Pt(1.4)
        dash(pn.line); pn.shadow.inherit = False
        text(sp, (x0 + x1) / 2, RY0 + 0.112, title, FS(6.3), edge, bold=True,
             w=x1 - x0 - 0.01)
        text(sp, (x0 + x1) / 2, RY0 + 0.042, sub, FS(fs_sub), "4D4D4D",
             w=x1 - x0 - 0.01)

    panel(0.006, 0.330, E_GEN, "deployment shifts   P1",
          "channel - multipath - fading - bandwidth - carrier/clock offset", 5.0)
    a1 = line(sp, 0.046, RY1 + 0.004, 0.046, sy0 - 0.006, E_GEN, 1.2, dashed=True)
    arrowhead(a1)
    panel(0.370, 0.650, RED, "rule-transfer test  (Sec. 4)",
          "3 curve statistics → log σepi", 5.8)
    a2 = line(sp, 0.398, RY1 + 0.004, 0.352, cy0 - 0.030, RED, 1.2, dashed=True)
    arrowhead(a2)
    a3 = line(sp, 0.612, yc - hh / 2 - 0.010, 0.578, RY1 + 0.004, RED, 1.2,
              dashed=True)
    arrowhead(a3)
    panel(0.690, 0.998, E_CP, "evaluation under P1  (Sec. 3, 5)",
          "coverage @ frozen Q  -  abstain", 5.6)
    a4 = line(sp, 0.875, yc - hh / 2 - 0.012, 0.865, RY1 + 0.004, E_CP, 1.2,
              dashed=True)
    arrowhead(a4)

    prs.save(OUT)
    print("[saved]", OUT)
    print("  슬라이드 %.2f x %.2f in, 도형 %d개 (전부 개별 편집 가능)"
          % (SW, SH, len(sl.shapes._spTree)))


if __name__ == "__main__":
    main()
